from pptx import Presentation
import pdfplumber


def extract_text_from_pptx(path):
    prs = Presentation(path)
    slides_text = []

    for i, slide in enumerate(prs.slides):
        title = slide.shapes.title.text.strip() if slide.shapes.title and slide.shapes.title.text else f"Slide {i + 1}"
        blocks = []

        for shape in slide.shapes:
            # Zones de texte
            if shape.has_text_frame:
                text = shape.text.strip()
                if text and text != title:
                    blocks.append(text)

            # Tableaux
            elif shape.has_table:
                table_text = []
                table = shape.table
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells]
                    table_text.append("\t".join(row_text))  # tabulation entre colonnes
                blocks.append("\n".join(table_text))  # retour à la ligne entre lignes

        slide_text = title + "\n\n" + "\n".join(blocks)
        slides_text.append(slide_text)

    return slides_text

def save_slides_to_txt(slides_text, output_path):
    with open(output_path, 'w', encoding='utf-8') as f:
        for i, slide_text in enumerate(slides_text):
            f.write(f"--- Slide {i + 1} ---\n{slide_text}\n\n")



def extract_text_from_pdf(path):
    slides_text = []

    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            blocks = []
            text = page.extract_text()

            # Supposons que la première ligne soit le "titre"
            lines = text.splitlines() if text else []
            #title = lines[0] if lines else f"Slide {i + 1}"
            body_lines = lines[1:] if len(lines) > 1 else []

            blocks.extend(body_lines)

            # Extraction des tableaux
            """
            try:
                tables = page.extract_tables()
                for table in tables:
                    row_strings = []
                    for row in table:
                        if row:
                            cleaned_row = [cell.strip().replace('\n', ' ') if cell else '' for cell in row]
                            row_strings.append(", ".join(cleaned_row))
                    flat_table = " / ".join(row_strings)
                    blocks.append(flat_table)
            except Exception as e:
                print(f"Erreur d'extraction de tableau sur la slide {i+1}: {e}")
            """
            # Nettoyage des blocs
            blocks = [block.strip() for block in blocks if block.strip()]
            slide_text = "\n".join(blocks)
            slides_text.append(slide_text)

    return slides_text



# Exemple d'utilisation
if __name__ == "__main__":
    input_path = "files/pg18.pdf"
    output_path = "files/pg18.txt"

    slides_text = extract_text_from_pdf(input_path)
    save_slides_to_txt(slides_text, output_path)

    print(f"Extraction PDF terminée. Résultat enregistré dans {output_path}")



